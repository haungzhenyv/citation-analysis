# -*- coding: utf-8 -*-
"""
make_intro_pdf.py —— 把「引用示例」中文介绍 Markdown 排成规范 PDF

用法：
    python make_intro_pdf.py <输入.md> <输出.pdf> [--title "文档总标题"]

输入 md 格式：每个条目以二级标题 `## 引用示例 N ...` 开头，其下是正文段落。

排版规范：
    总标题      黑体 16pt 加粗 蓝色 #1F4E79 居中
    条目标题    黑体 14pt 加粗 蓝色 #1F4E79 居中（「引用示例 N」）
    正文        宋体 12pt、两端对齐、首行缩进 2 字（24pt）、行距 1.12
    导出 PDF    用 Word COM（FileFormat=17 = wdFormatPDF）

依赖：python-docx、pywin32（Windows + 已安装 Word）
"""
import os
import re
import argparse

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

BLUE = RGBColor(0x1F, 0x4E, 0x79)
DEFAULT_TITLE = "论文被引分析 —— 大类一区期刊引用（中文介绍）"


def set_run(run, cn='宋体', en='Times New Roman', size=12, bold=False, color=None):
    """设置 run 的字体，显式指定 w:eastAsia 以免中文回退。"""
    run.font.name = en
    run.font.size = Pt(size)
    run.bold = bold
    run._element.rPr.rFonts.set(qn('w:eastAsia'), cn)
    if color is not None:
        run.font.color.rgb = color


def parse_items(text):
    """解析 `## 引用示例 N` 分节，返回 [(N, 正文), ...]"""
    pat = re.compile(r'^##\s*引用示例\s*(\d+)[^\n]*\n+(.*?)(?=\n##\s|\Z)', re.S | re.M)
    return [(int(m.group(1)), m.group(2).strip()) for m in pat.finditer(text)]


def build(md_path, docx_path, pdf_path, title):
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(2.2)
    sec.left_margin = Cm(2.6)
    sec.right_margin = Cm(2.6)

    st = doc.styles['Normal']
    st.font.name = 'Times New Roman'
    st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

    text = open(md_path, encoding='utf-8').read()
    items = parse_items(text)
    if not items:
        raise SystemExit("未找到任何 `## 引用示例 N` 条目，请检查输入 md 格式。")

    # 文档总标题
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(title)
    set_run(r, cn='黑体', size=16, bold=True, color=BLUE)

    # 逐条
    for n, body in items:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(f"引用示例 {n}")
        set_run(r, cn='黑体', size=14, bold=True, color=BLUE)

        for seg in [b.strip() for b in body.split('\n') if b.strip()]:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            pf = p.paragraph_format
            pf.space_before = Pt(0)
            pf.space_after = Pt(4)
            pf.line_spacing = 1.12
            pf.first_line_indent = Pt(24)
            r = p.add_run(seg)
            set_run(r, cn='宋体', size=12)

    doc.save(docx_path)
    print("DOCX SAVED:", docx_path, "items=", len(items))

    # Word COM 转 PDF
    import win32com.client
    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    try:
        word.DisplayAlerts = 0
    except Exception:
        pass
    d = word.Documents.Open(os.path.abspath(docx_path))
    d.SaveAs(os.path.abspath(pdf_path), FileFormat=17)  # wdFormatPDF
    d.Close(False)
    word.Quit()
    print("PDF SAVED:", pdf_path, os.path.exists(pdf_path))


def main():
    ap = argparse.ArgumentParser(description="把「引用示例」中文介绍 md 排成规范 PDF")
    ap.add_argument("md", help="输入 Markdown 路径")
    ap.add_argument("pdf", help="输出 PDF 路径（同名 .docx 会一并生成）")
    ap.add_argument("--title", default=DEFAULT_TITLE, help="文档总标题")
    a = ap.parse_args()
    docx_path = os.path.splitext(a.pdf)[0] + ".docx"
    build(a.md, docx_path, a.pdf, a.title)


if __name__ == "__main__":
    main()
